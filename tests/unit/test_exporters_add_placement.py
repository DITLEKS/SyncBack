"""Принятые добавления вставляются только по якорю; без якоря документ не меняется."""

import io

from docx import Document as DocxDocument

from app.domain.interfaces.document_exporter import AppliedChange
from app.infrastructure.exporters.docx_exporter import DocxExporter
from app.infrastructure.exporters.text_exporter import TextExporter

MD = (
    "# Сроки\n"
    "Срок ответа — 10 дней.\n"
    "Оплата картой.\n"
    "\n"
    "# Ответственность\n"
    "Стороны несут ответственность.\n"
)


def _add(new_text: str, *, anchor: str | None = None, section: str = "") -> AppliedChange:
    return AppliedChange(section_ref=section, change_type="add", old_text=anchor, new_text=new_text)


def _apply_text(changes: list[AppliedChange], text: str = MD) -> str:
    return TextExporter().apply_changes(text.encode("utf-8"), changes).decode("utf-8")


# ── txt / markdown ────────────────────────────────────────────────────────────────────────────────


def test_text_add_goes_right_after_anchor_line() -> None:
    result = _apply_text([_add("Срок может быть продлён.", anchor="Срок ответа — 10 дней.")])
    assert result == (
        "# Сроки\n"
        "Срок ответа — 10 дней.\n"
        "\n"
        "Срок может быть продлён.\n"
        "\n"
        "Оплата картой.\n"
        "\n"
        "# Ответственность\n"
        "Стороны несут ответственность.\n"
    )


def test_text_anchor_is_used_even_if_section_points_elsewhere() -> None:
    change = _add("Уточнение.", anchor="Стороны несут ответственность.", section="Сроки")
    result = _apply_text([change])
    assert result.endswith("Стороны несут ответственность.\n\nУточнение.\n")
    assert "Уточнение." not in result.split("# Ответственность")[0]


def test_text_add_without_anchor_is_not_applied() -> None:
    assert _apply_text([_add("Возможна оплата по счёту.", section="Сроки")]) == MD
    assert _apply_text([_add("Пробелы.", anchor="   ", section="Сроки")]) == MD


def test_text_add_with_unknown_anchor_is_not_applied() -> None:
    assert _apply_text([_add("Новый раздел.", anchor="такого текста нет", section="Сроки")]) == MD


def test_plain_txt_without_headings_uses_anchor() -> None:
    txt = "Первая строка.\nВторая строка.\n"
    assert _apply_text([_add("Вставка.", anchor="Первая строка.")], txt) == (
        "Первая строка.\n\nВставка.\n\nВторая строка.\n"
    )
    assert _apply_text([_add("Вставка.", section="Раздел")], txt) == txt


# ── docx ─────────────────────────────────────────────────────────────────────────────────────────────


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Сроки", style="Heading 1")
    document.add_paragraph("Срок ответа — 10 дней.")
    document.add_paragraph("Оплата картой.")
    document.add_paragraph("Ответственность", style="Heading 1")
    document.add_paragraph("Стороны несут ответственность.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


DOCX_PARAGRAPHS = [
    "Сроки",
    "Срок ответа — 10 дней.",
    "Оплата картой.",
    "Ответственность",
    "Стороны несут ответственность.",
]


def _paragraphs(raw: bytes) -> list[str]:
    return [p.text for p in DocxDocument(io.BytesIO(raw)).paragraphs]


def test_docx_add_goes_right_after_anchor_paragraph() -> None:
    raw = DocxExporter().apply_changes(
        _docx_bytes(), [_add("Срок может быть продлён.", anchor="Срок ответа — 10 дней.")]
    )
    assert _paragraphs(raw) == [
        "Сроки",
        "Срок ответа — 10 дней.",
        "Срок может быть продлён.",
        "Оплата картой.",
        "Ответственность",
        "Стороны несут ответственность.",
    ]


def test_docx_add_without_anchor_is_not_applied() -> None:
    raw = DocxExporter().apply_changes(
        _docx_bytes(), [_add("Возможна оплата по счёту.", section="Сроки")]
    )
    assert _paragraphs(raw) == DOCX_PARAGRAPHS


def test_docx_add_with_unknown_anchor_is_not_applied() -> None:
    raw = DocxExporter().apply_changes(
        _docx_bytes(), [_add("Новый раздел.", anchor="нет такого", section="Нет такого")]
    )
    assert _paragraphs(raw) == DOCX_PARAGRAPHS
