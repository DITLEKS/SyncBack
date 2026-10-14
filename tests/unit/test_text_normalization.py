"""Неразрывные пробелы в документе не мешают якорям совпадать с текстом.

Модель при копировании фрагмента возвращает обычные пробелы вместо U+00A0. Парсеры
нормализуют plain_text на входе, экспортёры — текст файла при сравнении, поэтому
маппер, «Чистовик» и экспорт видят одну и ту же строку.
"""

import io
import uuid

from docx import Document as DocxDocument

from app.domain.interfaces.document_exporter import AppliedChange
from app.domain.interfaces.llm_client import LLMSuggestionBatch, LLMSuggestionItem
from app.domain.text_normalization import normalize_text
from app.infrastructure.exporters.docx_exporter import DocxExporter
from app.infrastructure.exporters.text_exporter import TextExporter
from app.infrastructure.parsers.docx_parser import DocxParser
from app.infrastructure.parsers.markdown_parser import MarkdownParser
from app.infrastructure.parsers.txt_parser import TxtParser
from app.workers.pipeline.suggestion_mapper import map_to_suggestions

NBSP = "\u00a0"
RAW_LINE = f"Роль пользователя -{NBSP}customer,{NBSP}cook или courier."
CLEAN_LINE = "Роль пользователя - customer, cook или courier."
SSO = "Также поддерживается SSO."
ROLES = "Роли: customer, cook, courier."


def _docx_bytes(*paragraphs: str) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _paragraphs(raw: bytes) -> list[str]:
    return [p.text for p in DocxDocument(io.BytesIO(raw)).paragraphs]


def _change(change_type: str, old_text: str | None, new_text: str) -> AppliedChange:
    return AppliedChange(
        section_ref="Auth Service", change_type=change_type, old_text=old_text, new_text=new_text
    )


# ── normalize_text ───────────────────────────────────────────────────────────────────────────────


def test_normalize_replaces_space_like_and_removes_zero_width() -> None:
    assert normalize_text(RAW_LINE) == CLEAN_LINE
    assert normalize_text("a\u202fb\u2007c\u2009d") == "a b c d"
    assert normalize_text("\ufeffx\u200by") == "xy"


def test_normalize_applies_nfc_and_unix_newlines() -> None:
    assert normalize_text("е\u0308") == "ё"
    assert normalize_text("a\r\nb") == "a\nb"


def test_normalize_is_idempotent() -> None:
    once = normalize_text(RAW_LINE)
    assert normalize_text(once) == once


# ── парсеры ─────────────────────────────────────────────────────────────────────────────────────


def test_parsers_return_normalized_plain_text() -> None:
    assert TxtParser().parse(RAW_LINE.encode("utf-8")).plain_text == CLEAN_LINE
    md = MarkdownParser().parse(f"# Auth\n{RAW_LINE}\n".encode("utf-8"))
    assert md.plain_text == f"# Auth\n{CLEAN_LINE}\n"
    docx = DocxParser().parse(_docx_bytes("Auth Service", RAW_LINE))
    assert docx.plain_text == f"Auth Service\n{CLEAN_LINE}"


def test_markdown_section_offsets_are_computed_on_normalized_text() -> None:
    parsed = MarkdownParser().parse(f"\ufeff# Auth\n{RAW_LINE}\n".encode("utf-8"))
    section = parsed.sections[0]
    assert parsed.plain_text[section.start_offset : section.end_offset].startswith("# Auth")


# ── маппер ──────────────────────────────────────────────────────────────────────────────────────


def test_mapper_keeps_add_whose_anchor_has_plain_spaces() -> None:
    document_text = DocxParser().parse(_docx_bytes("Auth Service", RAW_LINE)).plain_text
    item = LLMSuggestionItem(
        section_ref="Auth Service",
        change_type="add",
        old_text=CLEAN_LINE,
        new_text=SSO,
        rationale=None,
        confidence_score=1.0,
    )
    batch = LLMSuggestionBatch(items=[item], raw_response="{}")
    result = map_to_suggestions(batch, uuid.uuid4(), uuid.uuid4(), document_text=document_text)
    assert len(result) == 1


# ── экспортёры ──────────────────────────────────────────────────────────────────────────────────


def test_text_exporter_matches_nbsp_in_file() -> None:
    raw = f"Auth Service\n{RAW_LINE}\nТаблица users\n".encode("utf-8")
    added = TextExporter().apply_changes(raw, [_change("add", CLEAN_LINE, SSO)])
    assert added.decode("utf-8") == (
        f"Auth Service\n{CLEAN_LINE}\n\n{SSO}\n\nТаблица users\n"
    )
    modified = TextExporter().apply_changes(raw, [_change("modify", CLEAN_LINE, ROLES)])
    assert modified.decode("utf-8") == f"Auth Service\n{ROLES}\nТаблица users\n"


def test_docx_exporter_matches_nbsp_in_paragraph() -> None:
    raw = _docx_bytes("Auth Service", RAW_LINE, "Таблица users")
    added = DocxExporter().apply_changes(raw, [_change("add", CLEAN_LINE, SSO)])
    assert _paragraphs(added) == ["Auth Service", RAW_LINE, SSO, "Таблица users"]
    modified = DocxExporter().apply_changes(raw, [_change("modify", CLEAN_LINE, ROLES)])
    assert _paragraphs(modified) == ["Auth Service", ROLES, "Таблица users"]
    deleted = DocxExporter().apply_changes(raw, [_change("delete", CLEAN_LINE, "")])
    assert _paragraphs(deleted) == ["Auth Service", "", "Таблица users"]
