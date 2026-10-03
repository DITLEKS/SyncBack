"""Документы проекта: загрузка, список, содержимое, скачивание, экспорт, источники."""

from __future__ import annotations

import io
import uuid

import pytest
from docx import Document as DocxDocument
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.infrastructure.db.models.audit_log import AuditLog
from tests.contract.conftest import create_project, register_and_login, upload_document

pytestmark = pytest.mark.asyncio


def _docx_bytes(*paragraphs: str) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_upload_list_get_document(client: AsyncClient) -> None:
    headers = await register_and_login(client)
    project_id = await create_project(client, headers)

    document = await upload_document(client, headers, project_id)
    assert document["name"] == "spec.txt"
    assert document["format"] == "txt"
    assert document["size_bytes"] == len(b"Hello world\n")
    assert document["status"] == "draft"

    response = await client.get(f"/api/v1/projects/{project_id}/documents", headers=headers)
    assert response.status_code == 200, response.text
    assert [d["id"] for d in response.json()["items"]] == [document["id"]]
    assert response.json()["total"] == 1

    response = await client.get(
        f"/api/v1/projects/{project_id}/documents/{document['id']}", headers=headers
    )
    assert response.status_code == 200, response.text
    assert response.json()["id"] == document["id"]

    response = await client.get(
        f"/api/v1/projects/{project_id}/documents/{uuid.uuid4()}", headers=headers
    )
    assert response.status_code == 404


async def test_upload_rejects_unsupported_format(client: AsyncClient) -> None:
    headers = await register_and_login(client)
    project_id = await create_project(client, headers)
    response = await client.post(
        f"/api/v1/projects/{project_id}/documents",
        files={"file": ("old.doc", b"x", "application/msword")},
        headers=headers,
    )
    assert response.status_code == 415


async def test_document_content_and_download_write_audit(
    client: AsyncClient, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    headers = await register_and_login(client)
    project_id = await create_project(client, headers)
    document = await upload_document(client, headers, project_id, content=b"Alpha\n\nBeta\n")
    base = f"/api/v1/projects/{project_id}/documents/{document['id']}"

    response = await client.get(f"{base}/content", headers=headers)
    assert response.status_code == 200, response.text
    assert "Alpha" in response.json()["plain_text"]

    response = await client.get(f"{base}/download", headers=headers)
    assert response.status_code == 200, response.text
    assert response.json()["download_url"].startswith("http://storage.test/")

    async with sessionmaker() as session:
        entries = (await session.execute(select(AuditLog))).scalars().all()
    assert [e.action for e in entries] == ["download"]
    assert str(entries[0].document_id) == document["id"]


async def test_export_in_source_format_only(client: AsyncClient) -> None:
    headers = await register_and_login(client)
    project_id = await create_project(client, headers)
    document = await upload_document(
        client,
        headers,
        project_id,
        filename="report.docx",
        content=_docx_bytes("Первый абзац", "Второй абзац"),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    base = f"/api/v1/projects/{project_id}/documents/{document['id']}/export"

    response = await client.get(base, headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"] == 'attachment; filename="report.docx"'
    assert DocxDocument(io.BytesIO(response.content)).paragraphs[0].text == "Первый абзац"

    # Конвертация между форматами не поддерживается — только исходный формат
    response = await client.get(base, params={"export_format": "md"}, headers=headers)
    assert response.status_code == 400

    response = await client.get(base, params={"export_format": "pdf"}, headers=headers)
    assert response.status_code == 400


async def test_export_text_document_applies_no_changes_when_none_accepted(
    client: AsyncClient,
) -> None:
    headers = await register_and_login(client)
    project_id = await create_project(client, headers)
    document = await upload_document(
        client, headers, project_id, filename="notes.md", content="# Заголовок\n\nТекст\n".encode()
    )
    response = await client.get(
        f"/api/v1/projects/{project_id}/documents/{document['id']}/export",
        params={"export_format": "md"},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"] == 'attachment; filename="notes.md"'
    assert response.headers["content-type"].startswith("text/markdown")
    assert response.text == "# Заголовок\n\nТекст\n"


async def test_attach_sources_to_document(client: AsyncClient) -> None:
    headers = await register_and_login(client)
    project_id = await create_project(client, headers)
    document = await upload_document(client, headers, project_id)

    response = await client.post(
        f"/api/v1/projects/{project_id}/sources/note",
        json={"name": "Глоссарий", "text_content": "термин — значение"},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    source_id = response.json()["id"]

    url = f"/api/v1/projects/{project_id}/documents/{document['id']}/sources"
    response = await client.post(url, json={"source_ids": [source_id]}, headers=headers)
    assert response.status_code == 200, response.text
    assert [s["id"] for s in response.json()["sources"]] == [source_id]
    assert response.json()["document"]["id"] == document["id"]

    # Повторное прикрепление идемпотентно
    response = await client.post(url, json={"source_ids": [source_id]}, headers=headers)
    assert response.status_code == 200, response.text

    response = await client.post(url, json={"source_ids": [str(uuid.uuid4())]}, headers=headers)
    assert response.status_code == 404


async def test_delete_document(client: AsyncClient) -> None:
    headers = await register_and_login(client)
    project_id = await create_project(client, headers)
    document = await upload_document(client, headers, project_id)
    url = f"/api/v1/projects/{project_id}/documents/{document['id']}"
    assert (await client.delete(url, headers=headers)).status_code == 204
    assert (await client.get(url, headers=headers)).status_code == 404
    assert (await client.delete(url, headers=headers)).status_code == 404


async def test_global_upload_requires_own_project(client: AsyncClient) -> None:
    headers = await register_and_login(client)
    project_id = await create_project(client, headers)

    response = await client.post(
        "/api/v1/documents",
        data={"project_id": project_id},
        files={"file": ("spec.txt", b"Hello", "text/plain")},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    assert response.json()["project_id"] == project_id

    response = await client.post(
        "/api/v1/documents",
        data={"project_id": project_id},
        files={"file": ("legacy.doc", b"x", "application/msword")},
        headers=headers,
    )
    assert response.status_code == 415, response.text

    other_headers = await register_and_login(client, email="other@example.com")
    response = await client.post(
        "/api/v1/documents",
        data={"project_id": project_id},
        files={"file": ("spec.txt", b"Hello", "text/plain")},
        headers=other_headers,
    )
    assert response.status_code == 404, response.text

    response = await client.get(f"/api/v1/projects/{project_id}/documents", headers=headers)
    assert response.json()["total"] == 1
